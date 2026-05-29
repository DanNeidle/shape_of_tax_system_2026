#!/usr/bin/env python3
import csv
import re
from pathlib import Path

from openpyxl import load_workbook


BASE = Path(__file__).resolve().parent
DATA_DIR = BASE.parent / "data" / "germany_tax"

EUROSTAT_NTL = DATA_DIR / "eurostat_national_tax_lists_2025_2026-04-22.xlsx"
BMF_CASH_TAXES = DATA_DIR / "bmf_steuereinnahmen_2024_calendar_year.xlsx"

EUROSTAT_OUTPUT = DATA_DIR / "germany_eurostat_tax_revenue_candidate_table_2024.csv"
BMF_OUTPUT = DATA_DIR / "germany_bmf_cash_tax_revenue_2024.csv"

GDP_EUR_M = 4_305_300

EUROSTAT_SOURCE = (
    "Eurostat, National Tax Lists - individual taxes, updated 22 Apr 2026; "
    "file: code/data/germany_tax/eurostat_national_tax_lists_2025_2026-04-22.xlsx"
)
BMF_SOURCE = (
    "Bundesministerium der Finanzen, Steuereinnahmen im 4. Quartal und Kalenderjahr 2024; "
    "file: code/data/germany_tax/bmf_steuereinnahmen_2024_calendar_year.xlsx"
)

SOCIAL_CONTRIBUTION_LEAF_CODES = {
    "D611C",
    "D612",
    "D613CE",
    "D613CS",
    "D613CN",
    "D613V",
}

UNCOLLECTIBLE_ADJUSTMENT_CODES = {
    "D995B",
    "D995C",
    "D995D",
    "D995E",
    "D995FE",
    "D995FS",
    "D995FN",
}

ECONOMIC_FUNCTION = {
    "C": "consumption",
    "AT": "alcohol/tobacco",
    "E": "energy",
    "P": "pollution",
    "T": "transport",
    "KS": "capital/stock of wealth",
    "KIH": "household capital income",
    "KIC": "corporate capital income",
    "KISE": "self-employed capital income",
    "LEYRS": "employer labour",
    "LEES": "employee labour",
    "LNON": "non-employed labour",
    "RP": "recurrent immovable property",
    "O": "other property",
    "_Z": "not classified",
}

BMF_TOTAL_LABELS = {
    "Gemeinschaftliche Steuern insgesamt",
    "Gewerbesteuerumlagen insgesamt",
    "Bundessteuern insgesamt",
    "Ländersteuern insgesamt",
    "Steuern insgesamt ohne Gemeindesteuern",
    "Gemeindesteuern Stadtstaaten insgesamt",
}


def clean_text(value):
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def normalise_key(value):
    value = clean_text(value).lower()
    value = value.replace("œ", "oe").replace("’", "'")
    value = re.sub(r"[^a-z0-9]+", " ", value)
    return re.sub(r"\s+", " ", value).strip()


def eur_and_pct_from_million(value_m):
    value_eur = round(float(value_m) * 1_000_000)
    pct = float(value_m) / GDP_EUR_M * 100
    return str(value_eur), f"{pct:.6f}"


def eurostat_explanation(sto, economic_function, property_environment_code):
    econ = ECONOMIC_FUNCTION.get(clean_text(economic_function), clean_text(economic_function))
    property_tax = ECONOMIC_FUNCTION.get(clean_text(property_environment_code), clean_text(property_environment_code))
    bits = [f"Eurostat National Tax Lists line classified as ESA code {sto}."]
    if econ and econ != "not classified":
        bits.append(f"Economic function: {econ}.")
    if property_tax and property_tax != "not classified":
        bits.append(f"Environmental/property tag: {property_tax}.")
    return " ".join(bits)


def read_eurostat_rows():
    wb = load_workbook(EUROSTAT_NTL, read_only=True, data_only=True)
    ws = wb["DE"]
    header = list(ws.iter_rows(min_row=8, max_row=8, values_only=True))[0]
    year_cols = {value: idx for idx, value in enumerate(header) if isinstance(value, int)}
    col_2024 = year_cols[2024]

    rows = []
    seen = set()
    current_parent_by_sto = {}
    for raw in ws.iter_rows(min_row=9, values_only=True):
        sto = clean_text(raw[0])
        detail = clean_text(raw[1])
        raw_tax_name = "" if raw[2] is None else str(raw[2])
        raw_tax_name_english = "" if raw[5] is None else str(raw[5])
        tax_name = clean_text(raw_tax_name)
        tax_name_english = clean_text(raw_tax_name_english)
        economic_function = clean_text(raw[9])
        property_environment_code = clean_text(raw[10])
        value_2024 = raw[col_2024]

        if not sto or not tax_name:
            continue
        if detail != "_T" and not raw_tax_name[:1].isspace() and not isinstance(value_2024, (int, float)):
            current_parent_by_sto[sto] = (tax_name, tax_name_english)
        if raw_tax_name[:1].isspace() and sto in current_parent_by_sto:
            parent_name, parent_name_english = current_parent_by_sto[sto]
            tax_name = f"{parent_name} - {tax_name}"
            tax_name_english = f"{parent_name_english} - {tax_name_english}"

        keep = False
        is_uncollectible_adjustment = False
        note_parts = []
        if detail != "_T":
            keep = True
        elif sto in SOCIAL_CONTRIBUTION_LEAF_CODES:
            keep = True
            note_parts.append("Social-contribution component retained where Eurostat does not provide a lower national-tax-list split.")
        elif sto in UNCOLLECTIBLE_ADJUSTMENT_CODES:
            keep = True
            is_uncollectible_adjustment = True
            note_parts.append("Negative D995 adjustment for taxes/social contributions assessed but unlikely to be collected.")

        if not keep or not isinstance(value_2024, (int, float)):
            continue
        if value_2024 <= 0 and not is_uncollectible_adjustment:
            continue
        if is_uncollectible_adjustment:
            value_2024 = -abs(value_2024)
            tax_name = f"Less: {tax_name}"
            tax_name_english = f"Less: {tax_name_english}"

        key = (normalise_key(tax_name), normalise_key(tax_name_english), round(float(value_2024), 6))
        if key in seen:
            continue
        seen.add(key)

        revenue_eur, revenue_pct_gdp = eur_and_pct_from_million(value_2024)
        rows.append({
            "tax_name": tax_name,
            "tax_name_english": tax_name_english,
            "english_explanation": eurostat_explanation(sto, economic_function, property_environment_code),
            "revenue_eur": revenue_eur,
            "revenue_pct_gdp": revenue_pct_gdp,
            "source_for_data": f"{EUROSTAT_SOURCE}; DE sheet; STO {sto}; DETAILS {detail}; 2024",
            "notes": "Eurostat NTL covers general government plus EU institutions (S.13 + S.212), so local taxes and compulsory social contributions are in scope.",
        })
    return rows


def bmf_row_name(raw):
    for idx in (4, 3, 2, 1):
        value = clean_text(raw[idx])
        if value and value != "davon:" and value != "+":
            return value
    return ""


def read_bmf_cash_rows():
    wb = load_workbook(BMF_CASH_TAXES, read_only=True, data_only=True)
    rows = []
    for sheet_name in ["1 Steuerarten", "3u4 weitere Angaben"]:
        ws = wb[sheet_name]
        current_group = ""
        for raw in ws.iter_rows(values_only=True):
            values = list(raw) + [None] * 12
            if clean_text(values[1]) and not isinstance(values[8], (int, float)):
                current_group = clean_text(values[1])
            name = bmf_row_name(values)
            value_2024_thousand = values[8]
            if not name or not isinstance(value_2024_thousand, (int, float)):
                continue
            is_total = name in BMF_TOTAL_LABELS or name.endswith("insgesamt")
            rows.append({
                "tax_name": name,
                "group": current_group,
                "revenue_eur": str(round(float(value_2024_thousand) * 1_000)),
                "source_for_data": f"{BMF_SOURCE}; sheet: {sheet_name}; annual 2024 column; source values in thousand euros",
                "is_total_or_memo": "yes" if is_total else "no",
                "notes": "BMF cash tax revenue table. Main sheet excludes municipal taxes; sheet 3 includes memo rows for city-state municipal taxes and energy-tax sub-splits.",
            })
    return rows


def write_csv(path, rows, fieldnames):
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def main():
    eurostat_rows = read_eurostat_rows()
    bmf_rows = read_bmf_cash_rows()
    write_csv(
        EUROSTAT_OUTPUT,
        eurostat_rows,
        ["tax_name", "tax_name_english", "english_explanation", "revenue_eur", "revenue_pct_gdp", "source_for_data", "notes"],
    )
    write_csv(
        BMF_OUTPUT,
        bmf_rows,
        ["tax_name", "group", "revenue_eur", "source_for_data", "is_total_or_memo", "notes"],
    )

    eurostat_total = sum(int(row["revenue_eur"]) for row in eurostat_rows)
    bmf_total = next(
        (
            int(row["revenue_eur"])
            for row in bmf_rows
            if row["tax_name"] == "Steuern insgesamt ohne Gemeindesteuern"
        ),
        None,
    )
    print(f"Wrote {EUROSTAT_OUTPUT}")
    print(f"Eurostat candidate rows: {len(eurostat_rows)}")
    print(f"Eurostat candidate signed total: EUR {eurostat_total:,}")
    print(f"Eurostat candidate signed total % GDP: {eurostat_total / (GDP_EUR_M * 1_000_000) * 100:.6f}%")
    print(f"Wrote {BMF_OUTPUT}")
    print(f"BMF cash rows: {len(bmf_rows)}")
    print(f"BMF total excluding municipal taxes: EUR {bmf_total:,}" if bmf_total else "BMF total excluding municipal taxes: not found")
    print("BMF rows are not additive: the table includes totals, subtotals and subcomponents.")


if __name__ == "__main__":
    main()
