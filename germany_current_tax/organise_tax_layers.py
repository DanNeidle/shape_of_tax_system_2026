#!/usr/bin/env python3
import csv
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


BASE = Path(__file__).resolve().parent
DATA_DIR = BASE.parent / "data" / "germany_tax"

EUROSTAT_INPUT = DATA_DIR / "germany_eurostat_tax_revenue_candidate_table_2024.csv"
BMF_INPUT = DATA_DIR / "germany_bmf_cash_tax_revenue_2024.csv"
BMF_WORKBOOK = DATA_DIR / "bmf_steuereinnahmen_2024_calendar_year.xlsx"
OUTPUT = DATA_DIR / "germany_tax_revenue_layered_2024.csv"

GDP_EUR = 4_305_300_000_000

FIELDNAMES = [
    "tax_name",
    "tax_name_english",
    "english_explanation",
    "revenue_eur",
    "revenue_pct_gdp",
    "source_for_data",
    "notes",
    "chart_label_en",
    "layer_1",
    "layer_2",
    "layer_3",
    "layer_4",
    "classification_notes",
    "display_explanation",
    "chart_include",
    "chart_revenue_eur",
    "chart_revenue_pct_gdp",
    "netted_adjustments_eur",
    "netting_notes",
]


SIMPLE_RULES = {
    "Zölle": {
        "label": "Customs duties",
        "english": "Customs duties",
        "layers": ("Goods/services", "Import duties", "Customs duties", ""),
        "classification": "Direct UK equivalent: customs duties.",
        "explanation": "Customs duties on imports.",
    },
    "Verbrauchsteuern auf Einfuhren": {
        "label": "Import excise duties",
        "english": "Import excise duties",
        "layers": ("Goods/services", "Import duties", "Import excise duties", ""),
        "classification": "Import-related excise duty.",
        "explanation": "Excise duties on imported goods.",
    },
    "Stromsteuer": {
        "label": "Electricity tax",
        "english": "Electricity tax",
        "layers": ("Goods/services", "Energy and fuel duties", "Energy/fuel consumption duties", ""),
        "classification": "UK-equivalent treatment: energy/fuel consumption duties sit with goods/services.",
        "explanation": "Tax on electricity consumption.",
    },
    "Mineralölsteuer / Energiesteuer ab 2006": {
        "label": "Energy tax",
        "english": "Energy tax",
        "layers": ("Goods/services", "Energy and fuel duties", "Energy/fuel consumption duties", ""),
        "classification": "UK-equivalent treatment: energy/fuel consumption duties sit with goods/services.",
        "explanation": "Excise duty on fuel and other energy products.",
    },
    "Tabaksteuer": {
        "label": "Tobacco duty",
        "english": "Tobacco duty",
        "layers": ("Goods/services", "Alcohol, tobacco and drinks", "Excise duties", ""),
        "classification": "Direct UK equivalent: tobacco duty.",
        "explanation": "Excise duty on tobacco products.",
    },
    "Kaffeesteuer": {
        "label": "Coffee duty",
        "english": "Coffee duty",
        "layers": ("Goods/services", "Alcohol, tobacco and drinks", "Excise duties", ""),
        "classification": "German-specific product duty.",
        "explanation": "Excise duty on coffee.",
    },
    "Schaumweinsteuer": {
        "label": "Sparkling wine duty",
        "english": "Sparkling wine duty",
        "layers": ("Goods/services", "Alcohol, tobacco and drinks", "Excise duties", ""),
        "classification": "Alcohol duty.",
        "explanation": "Excise duty on sparkling wine.",
    },
    "Biersteuer": {
        "label": "Beer duty",
        "english": "Beer duty",
        "layers": ("Goods/services", "Alcohol, tobacco and drinks", "Excise duties", ""),
        "classification": "Direct UK equivalent: beer duty.",
        "explanation": "Excise duty on beer.",
    },
    "Förderzahlungen nach KWKG": {
        "label": "CHP levy",
        "english": "CHP levy",
        "layers": ("Environmental", "Energy-transition levies", "Electricity-system levies", ""),
        "classification": "Energy-transition levy treated as an environmental/energy item.",
        "explanation": "Levy used to fund combined heat and power support.",
    },
    "Förderzahlungen Offshore-Netzumlage": {
        "label": "Offshore grid levy",
        "english": "Offshore grid levy",
        "layers": ("Environmental", "Energy-transition levies", "Electricity-system levies", ""),
        "classification": "Energy-transition levy treated as an environmental/energy item.",
        "explanation": "Levy used to fund offshore electricity-grid costs.",
    },
    "Grunderwerbsteuer": {
        "label": "Property transfer tax",
        "english": "Property transfer tax",
        "layers": ("Land", "Property transaction taxes", "Property transfer taxes", ""),
        "classification": "Direct UK equivalent: stamp duty on property transactions.",
        "explanation": "Tax on purchases of land and buildings.",
    },
    "Sonstige Vergnügungsteuer": {
        "label": "Entertainment tax",
        "english": "Entertainment tax",
        "layers": ("Goods/services", "Betting, gaming and entertainment", "Entertainment taxes", ""),
        "classification": "Local consumption tax on entertainment.",
        "explanation": "Local tax on entertainment activities.",
    },
    "Rennwett-Lotteriesteuer": {
        "label": "Betting/lottery tax",
        "english": "Betting and lottery tax",
        "layers": ("Goods/services", "Betting, gaming and entertainment", "Betting and lottery taxes", ""),
        "classification": "Direct UK equivalent: betting and lottery duties.",
        "explanation": "Tax on betting and lottery activities.",
    },
    "Spielbankabgabe": {
        "label": "Casino levy",
        "english": "Casino levy",
        "layers": ("Goods/services", "Betting, gaming and entertainment", "Casino taxes", ""),
        "classification": "Gaming levy.",
        "explanation": "Tax or levy on casino gaming.",
    },
    "Versicherungsteuer": {
        "label": "Insurance tax",
        "english": "Insurance tax",
        "layers": ("Goods/services", "Insurance taxes", "Insurance-premium taxes", ""),
        "classification": "Direct UK equivalent: insurance premium tax.",
        "explanation": "Tax on insurance premiums or contracts.",
    },
    "Feuerschutzsteuer": {
        "label": "Fire protection tax",
        "english": "Fire protection tax",
        "layers": ("Goods/services", "Insurance taxes", "Insurance-premium taxes", ""),
        "classification": "Insurance-related tax.",
        "explanation": "Tax on fire insurance used to fund fire protection.",
    },
    "Luftverkehrsteuer": {
        "label": "Air travel tax",
        "english": "Air travel tax",
        "layers": ("Goods/services", "Transport and vehicle taxes", "Transport taxes", ""),
        "classification": "Direct UK equivalent: air passenger duty.",
        "explanation": "Tax on departing airline passengers.",
    },
    "Beitrag zum Erdölbevorratungsverband": {
        "label": "Oil stockpiling levy",
        "english": "Oil stockpiling levy",
        "layers": ("Goods/services", "Energy and fuel duties", "Strategic stock levies", ""),
        "classification": "Energy-security levy.",
        "explanation": "Levy used to fund strategic oil stocks.",
    },
    "übrige Gemeindesteuern": {
        "label": "Other municipal taxes",
        "english": "Other municipal taxes",
        "layers": ("Other", "Small local taxes", "Other municipal taxes", ""),
        "classification": "Small residual local taxes retained because Eurostat records them as tax revenue.",
        "explanation": "Other small municipal taxes not separately identified in the source table.",
    },
    "Emissionsberechtigungen": {
        "label": "Emissions allowances",
        "english": "Emissions allowances",
        "layers": ("Environmental", "Carbon and emissions taxes", "Emissions trading", ""),
        "classification": "UK-equivalent treatment: emissions-auction receipts sit with environmental taxes.",
        "explanation": "Revenue from auctioning emissions allowances.",
    },
    "Nationale CO2 - Abgabe": {
        "label": "National CO2 levy",
        "english": "National CO2 levy",
        "layers": ("Environmental", "Carbon and emissions taxes", "Carbon levies", ""),
        "classification": "Carbon-pricing levy.",
        "explanation": "National carbon levy on fossil fuels.",
    },
    "Steuerähnliche Einnahmen": {
        "label": "Tax-like receipts",
        "english": "Tax-like receipts",
        "layers": ("Other", "Tax-like fees and charges", "Tax-like receipts", ""),
        "classification": "Included because Eurostat records the receipts as taxes in the national accounts.",
        "explanation": "Tax-like receipts recorded as production taxes in the national accounts.",
    },
    "Übrige Produktionsabgaben": {
        "label": "Other production levies",
        "english": "Other production levies",
        "layers": ("Business", "Business production taxes", "Other production levies", ""),
        "classification": "Business production tax.",
        "explanation": "Other small taxes on production.",
    },
    "Bankenabgabe": {
        "label": "Bank levy",
        "english": "Bank levy",
        "layers": ("Business", "Financial-sector levies", "Bank levies", ""),
        "classification": "Direct UK equivalent: bank levy.",
        "explanation": "Levy on banks.",
    },
    "Beitrag zum Einlagensicherungsfond": {
        "label": "Deposit protection levy",
        "english": "Deposit protection levy",
        "layers": ("Business", "Financial-sector levies", "Bank levies", ""),
        "classification": "Financial-sector levy.",
        "explanation": "Banking-sector levy for deposit protection.",
    },
    "Lohnsteuer": {
        "label": "Wage tax",
        "english": "Wage tax",
        "layers": ("Employment", "Personal income taxes", "Employment income taxes", ""),
        "classification": "Direct UK equivalent: PAYE-style withholding on wages.",
        "explanation": "Payroll withholding on employment income.",
    },
    "Veranlagte Einkommensteuer": {
        "label": "Assessed income tax",
        "english": "Assessed income tax",
        "layers": ("Employment", "Personal income taxes", "Self-assessed income taxes", ""),
        "classification": "Personal income tax assessed through returns.",
        "explanation": "Income tax assessed through tax returns.",
    },
    "Körperschaftsteuer": {
        "label": "Corporation tax",
        "english": "Corporation tax",
        "layers": ("Business", "Profit and sector taxes", "Corporation tax", ""),
        "classification": "Direct UK equivalent: corporation tax.",
        "explanation": "Tax on company profits.",
    },
    "EU-Energiekrisenbeitrag": {
        "label": "Energy crisis levy",
        "english": "EU energy crisis contribution",
        "layers": ("Business", "Profit and sector taxes", "Windfall and crisis levies", ""),
        "classification": "Temporary sectoral profit levy.",
        "explanation": "Temporary levy on energy-sector windfall profits.",
    },
    "Einkommensteuer der übrigen Welt": {
        "label": "Non-resident income tax",
        "english": "Non-resident income tax",
        "layers": ("Employment", "Personal income taxes", "Non-resident income taxes", ""),
        "classification": "Income-tax line attributed to the rest of the world in the national accounts.",
        "explanation": "Income tax attributed to non-residents in the national accounts.",
    },
    "Hundesteuer": {
        "label": "Dog tax",
        "english": "Dog tax",
        "layers": ("Other", "Small local taxes", "Household local taxes", ""),
        "classification": "German-specific local tax.",
        "explanation": "Local tax on dog ownership.",
    },
    "Jagd- und Fischereisteuer": {
        "label": "Hunt/fishing tax",
        "english": "Hunting and fishing tax",
        "layers": ("Other", "Small local taxes", "Household local taxes", ""),
        "classification": "Local licence-related tax.",
        "explanation": "Local hunting and fishing tax.",
    },
    "Erbschaftsteuer": {
        "label": "Inheritance tax",
        "english": "Inheritance tax",
        "layers": ("Wealth", "Inheritance and gifts", "Inheritance and gift taxes", ""),
        "classification": "Direct UK equivalent: inheritance tax.",
        "explanation": "Tax on inheritances and lifetime gifts.",
    },
    "Tatsächliche Pflichtsozialbeiträge der Arbeitgeber": {
        "label": "Employer social contribs",
        "english": "Employer social contributions",
        "layers": ("Employment", "Social contributions", "Employer contributions", ""),
        "classification": "Compulsory social contribution.",
        "explanation": "Compulsory employer social security contributions on employees' pay.",
    },
    "unterstellte Sozialbeiträge": {
        "label": "Imputed employer contribs",
        "english": "Imputed employer social contributions",
        "layers": ("Employment", "Social contributions", "Employer contributions", ""),
        "classification": "National-accounts social-contribution imputation.",
        "explanation": "Social benefits paid directly by employers and recorded as employer social contributions in the national accounts.",
    },
    "Tatsächliche Pflichtsozialbeiträge der Arbeitnehmer": {
        "label": "Employee social contribs",
        "english": "Employee social contributions",
        "layers": ("Employment", "Social contributions", "Employee and household contributions", ""),
        "classification": "Compulsory social contribution.",
        "explanation": "Compulsory employee social security contributions deducted from wages.",
    },
    "Tatsächliche Pflichtsozialbeiträge der Selbständigen": {
        "label": "Self-employed contribs",
        "english": "Self-employed social contributions",
        "layers": ("Employment", "Social contributions", "Employee and household contributions", ""),
        "classification": "Compulsory social contribution.",
        "explanation": "Compulsory social security contributions paid by self-employed workers.",
    },
    "Tatsächliche Pflichtsozialbeiträge der Nichterwerbstätigen": {
        "label": "Non-employed contribs",
        "english": "Non-employed social contributions",
        "layers": ("Employment", "Social contributions", "Employee and household contributions", ""),
        "classification": "Compulsory social contribution.",
        "explanation": "Compulsory social security contributions paid by people who are not in employment.",
    },
    "Tatsächliche freiwillige Sozialbeiträge der privaten Haushalte": {
        "label": "Voluntary social contribs",
        "english": "Voluntary social contributions",
        "layers": ("Employment", "Social contributions", "Employee and household contributions", ""),
        "classification": "Voluntary household social contribution included by Eurostat in the social-contribution total.",
        "explanation": "Voluntary social security contributions paid by households.",
    },
}

GROUP_RULES = {
    "VAT": {
        "source_names": [
            "Umsatzsteuer (abzüglich Unterkompensation)",
            "Einfuhrumsatzsteuer",
            "Unterkompensation Umsatzsteuer",
        ],
        "label": "VAT",
        "english": "Value added tax",
        "layers": ("Goods/services", "VAT", "VAT", ""),
        "classification": "Direct UK equivalent: VAT. The Eurostat control includes import VAT and a VAT under-compensation adjustment.",
        "explanation": "Germany's value added tax, including import VAT and the national-accounts VAT adjustment.",
    },
    "Property tax": {
        "source_names": ["Grundsteuer A", "Grundsteuer B"],
        "label": "Property tax",
        "english": "Property tax",
        "layers": ("Land", "Recurrent land and property taxes", "Property taxes", ""),
        "classification": "Direct UK equivalent: recurrent local taxes on land and buildings.",
        "explanation": "Local recurrent tax on land and buildings.",
    },
    "Motor vehicle tax": {
        "source_names": ["Kfz-Steuer von Unternehmen", "Kraftfahrzeugsteuer von privaten Haushalten"],
        "label": "Motor vehicle tax",
        "english": "Motor vehicle tax",
        "layers": ("Goods/services", "Transport and vehicle taxes", "Vehicle taxes", ""),
        "classification": "Direct UK equivalent: vehicle excise duty-style tax.",
        "explanation": "Tax on owning or using motor vehicles.",
    },
    "Administrative fees": {
        "source_names": ["Verwaltungsgebühren von Unternehmen", "Verwaltungsgebühren von privaten Haushalten"],
        "label": "Administrative fees",
        "english": "Administrative fees",
        "layers": ("Other", "Tax-like fees and charges", "Administrative fees", ""),
        "classification": "Included because Eurostat records the receipts as taxes in the national accounts.",
        "explanation": "Administrative charges treated as tax revenue in the national accounts.",
    },
    "Broadcasting contribution": {
        "source_names": ["Rundfunkbeitrag der Unternehmen", "Rundfunkbeitrag der privaten Haushalte"],
        "label": "Broadcasting charge",
        "english": "Broadcasting contribution",
        "layers": ("Other", "Tax-like fees and charges", "Broadcasting contribution", ""),
        "classification": "Compulsory contribution recorded as tax revenue by Eurostat.",
        "explanation": "Compulsory broadcasting contribution paid by households and businesses.",
    },
    "Trade tax": {
        "source_names": ["Gewerbesteuer PHH", "Gewerbesteuer KG"],
        "label": "Trade tax",
        "english": "Trade tax",
        "layers": ("Business", "Profit and sector taxes", "Local business profit taxes", ""),
        "classification": "Local business tax on profits.",
        "explanation": "Local business tax on business profits.",
    },
    "Holding gains taxes": {
        "source_names": ["Steuern auf Umbewertungsgewinne"],
        "label": "Holding gains taxes",
        "english": "Taxes on holding gains",
        "layers": ("Wealth", "Financial capital taxes", "Holding gains taxes", ""),
        "classification": "National-accounts tax line on holding or revaluation gains.",
        "explanation": "Taxes on holding gains or revaluation gains recorded in the national accounts.",
    },
}

BMF_PROPORTION_SPLITS = [
    {
        "source_names": ["Alkoholsteuer / Branntweinabgaben / einschl. Alcopops ab 3. Quartal 2004"],
        "components": [
            ("Alkoholsteuer (bis 2017 Branntweinsteuer)", "Alcohol duty", "Alcohol duty", "Excise duty on spirits and other alcohol.", "Goods/services", "Alcohol, tobacco and drinks", "Excise duties"),
            ("Zwischenerzeugnissteuer", "Intermediate alcohol", "Intermediate alcohol duty", "Excise duty on intermediate alcoholic products.", "Goods/services", "Alcohol, tobacco and drinks", "Excise duties"),
            ("Alkopopsteuer", "Alcopop duty", "Alcopop duty", "Excise duty on alcopop drinks.", "Goods/services", "Alcohol, tobacco and drinks", "Excise duties"),
        ],
        "classification": "Eurostat groups these alcohol taxes; BMF 2024 cash receipts are used to split the Eurostat control total.",
    },
    {
        "source_names": ["sonstige Verbrauchsteuern"],
        "components": [
            ("Pauschalierte Einfuhrabgabe", "Flat-rate import charge", "Flat-rate import charge", "Flat-rate import charge.", "Goods/services", "Import duties", "Other import charges"),
            ("sonstige Bundessteuern", "Other federal taxes", "Other federal taxes", "Other very small federal taxes.", "Other", "Other small taxes", "Other federal taxes"),
        ],
        "classification": "Eurostat groups these very small taxes; BMF 2024 cash receipts are used to split the Eurostat control total.",
    },
    {
        "source_names": ["Kapitalertragsteuer und Zinsabschlag"],
        "components": [
            ("nicht veranlagte Steuern vom Ertrag", "Investment withholding", "Investment income withholding tax", "Withholding tax on investment income not assessed through tax returns.", "Wealth", "Financial capital taxes", "Investment income taxes"),
            ("Abgeltungsteuer auf Zins- und Veräußerungserträge", "Final withholding tax", "Final withholding tax on interest and capital gains", "Flat withholding tax on interest and capital gains.", "Wealth", "Financial capital taxes", "Investment income taxes"),
        ],
        "classification": "Eurostat gives sectoral capital-income-tax rows; BMF 2024 cash receipts are used to split them into statutory tax names.",
    },
]


def read_csv(path):
    with path.open(encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def read_eurostat():
    rows = read_csv(EUROSTAT_INPUT)
    by_name = defaultdict(list)
    for row in rows:
        by_name[row["tax_name"]].append(row)
    return rows, by_name


def read_bmf():
    rows = read_csv(BMF_INPUT)
    by_name = defaultdict(list)
    for row in rows:
        by_name[row["tax_name"]].append(row)
    return rows, by_name


def bmf_amount(by_name, tax_name):
    rows = [
        row for row in by_name[tax_name]
        if row["is_total_or_memo"] == "no" and int(row["revenue_eur"]) > 0
    ]
    main_sheet_rows = [row for row in rows if "sheet: 1 Steuerarten" in row["source_for_data"]]
    if main_sheet_rows:
        rows = main_sheet_rows
    if not rows:
        raise KeyError(f"Missing positive BMF row for {tax_name}")
    if len(rows) > 1:
        values = sorted({int(row["revenue_eur"]) for row in rows})
        if len(values) == 1:
            return values[0]
        raise ValueError(f"Ambiguous BMF rows for {tax_name}: {values}")
    return int(rows[0]["revenue_eur"])


def solidarity_components():
    wb = load_workbook(BMF_WORKBOOK, read_only=True, data_only=True)
    ws = wb["3u4 weitere Angaben"]
    rows = list(ws.iter_rows(values_only=True))
    wanted_rows = {
        "wage": 18,
        "assessed": 20,
        "capital_non_assessed": 21,
        "corporation": 22,
        "capital_final": 23,
        "total": 24,
    }
    values = {
        key: round(float(rows[row_number - 1][8]) * 1_000)
        for key, row_number in wanted_rows.items()
    }
    component_total = (
        values["wage"]
        + values["assessed"]
        + values["capital_non_assessed"]
        + values["corporation"]
        + values["capital_final"]
    )
    rounding_delta = values["total"] - component_total
    values["capital_final"] += rounding_delta
    component_total += rounding_delta
    if component_total != values["total"]:
        raise ValueError(f"Solidarity surcharge split mismatch: {component_total} vs {values['total']}")
    return values


def source_rows(by_name, names):
    rows = []
    for name in names:
        if name not in by_name:
            raise KeyError(f"Missing Eurostat row: {name}")
        rows.extend(by_name[name])
    return rows


def row_total(rows):
    return sum(int(row["revenue_eur"]) for row in rows)


def pct(value):
    return f"{value / GDP_EUR * 100:.6f}"


def joined(values):
    unique = []
    for value in values:
        if value and value not in unique:
            unique.append(value)
    return "; ".join(unique)


def make_entry(
    *,
    source,
    label,
    english,
    explanation,
    layers,
    classification,
    chart_value=None,
    netted=0,
    netting_notes="",
    extra_source="",
    extra_notes="",
):
    raw_value = row_total(source)
    if chart_value is None:
        chart_value = raw_value + netted
    tax_name = joined(row["tax_name"] for row in source)
    source_text = joined(row["source_for_data"] for row in source)
    if extra_source:
        source_text = joined([source_text, extra_source])
    notes = joined(row["notes"] for row in source)
    if extra_notes:
        notes = joined([notes, extra_notes])
    layer_1, layer_2, layer_3, layer_4 = (*layers, "")[:4]
    return {
        "tax_name": tax_name,
        "tax_name_english": english,
        "english_explanation": explanation,
        "revenue_eur": str(raw_value),
        "revenue_pct_gdp": pct(raw_value),
        "source_for_data": source_text,
        "notes": notes,
        "chart_label_en": label,
        "layer_1": layer_1,
        "layer_2": layer_2,
        "layer_3": layer_3,
        "layer_4": layer_4,
        "classification_notes": classification,
        "display_explanation": explanation,
        "chart_include": "yes" if chart_value > 0 else "no",
        "chart_revenue_eur": str(chart_value) if chart_value > 0 else "",
        "chart_revenue_pct_gdp": pct(chart_value) if chart_value > 0 else "",
        "netted_adjustments_eur": str(netted),
        "netting_notes": netting_notes,
    }


def make_bmf_split_entries(euro_by_name, bmf_by_name):
    rows = []
    bmf_source = "Bundesministerium der Finanzen, 2024 cash tax receipts, used only to split an equivalent Eurostat control line."
    for split in BMF_PROPORTION_SPLITS:
        source = source_rows(euro_by_name, split["source_names"])
        total = row_total(source)
        bmf_values = [bmf_amount(bmf_by_name, component[0]) for component in split["components"]]
        bmf_total = sum(bmf_values)
        allocated = []
        running = 0
        for idx, value in enumerate(bmf_values):
            if idx == len(bmf_values) - 1:
                amount = total - running
            else:
                amount = round(total * value / bmf_total)
                running += amount
            allocated.append(amount)
        for component, amount, bmf_value in zip(split["components"], allocated, bmf_values):
            german, label, english, explanation, layer_1, layer_2, layer_3 = component
            rows.append(make_entry(
                source=source,
                label=label,
                english=english,
                explanation=explanation,
                layers=(layer_1, layer_2, layer_3, german),
                classification=split["classification"],
                chart_value=amount,
                extra_source=bmf_source,
                extra_notes=f"Allocated from the Eurostat row using BMF cash receipt share for {german}: EUR {bmf_value:,}.",
            ))
    return rows


def make_simple_entries(euro_by_name, solidarity):
    rows = []
    used_names = set()

    def add(names, rule, netted=0, netting_notes=""):
        nonlocal rows
        source = source_rows(euro_by_name, names)
        rows.append(make_entry(
            source=source,
            label=rule["label"],
            english=rule["english"],
            explanation=rule["explanation"],
            layers=rule["layers"],
            classification=rule["classification"],
            netted=netted,
            netting_notes=netting_notes,
        ))
        used_names.update(names)

    for rule in GROUP_RULES.values():
        add(rule["source_names"], rule)

    for name, rule in SIMPLE_RULES.items():
        adjustment = 0
        notes = ""
        if name == "Lohnsteuer":
            adjustment = -solidarity["wage"]
            notes = "BMF cash receipts identify this amount as the solidarity surcharge on wage tax; it is shown separately."
        elif name == "Veranlagte Einkommensteuer":
            adjustment = -solidarity["assessed"]
            notes = "BMF cash receipts identify this amount as the solidarity surcharge on assessed income tax; it is shown separately."
        elif name == "Körperschaftsteuer":
            adjustment = -solidarity["corporation"]
            notes = "BMF cash receipts identify this amount as the solidarity surcharge on corporation tax; it is shown separately."
        add([name], rule, adjustment, notes)

    split_names = {name for split in BMF_PROPORTION_SPLITS for name in split["source_names"]}
    used_names.update(split_names)
    missing = sorted(set(euro_by_name) - used_names)
    if missing:
        raise ValueError(f"Eurostat rows not assigned to the layered table: {missing}")
    return rows


def make_solidarity_entry(euro_by_name, solidarity):
    related_sources = source_rows(euro_by_name, [
        "Lohnsteuer",
        "Veranlagte Einkommensteuer",
        "Kapitalertragsteuer und Zinsabschlag",
        "Körperschaftsteuer",
    ])
    return make_entry(
        source=related_sources,
        label="Solidarity surcharge",
        english="Solidarity surcharge",
        explanation="Surcharge on income tax, capital income tax and corporation tax.",
        layers=("Other", "Income-tax surcharges", "Solidarity surcharge", ""),
        classification="Statutory tax shown separately using BMF 2024 cash receipts; Eurostat does not list it as a separate Germany NTL row.",
        chart_value=solidarity["total"],
        extra_source="Bundesministerium der Finanzen, 2024 cash tax receipts; sheet 3u4 weitere Angaben; Solidaritätszuschlag insgesamt.",
        extra_notes="The same amount is netted out of the relevant Eurostat income, capital-income and corporation-tax lines so the chart still reconciles to the Eurostat control total.",
    )


def apply_capital_solidarity_adjustment(rows, solidarity):
    capital_adjustment = solidarity["capital_non_assessed"] + solidarity["capital_final"]
    capital_labels = {"Investment withholding", "Final withholding tax"}
    capital_total = sum(int(row["chart_revenue_eur"]) for row in rows if row["chart_label_en"] in capital_labels)
    running = 0
    for row in rows:
        if row["chart_label_en"] not in capital_labels:
            continue
        chart_value = int(row["chart_revenue_eur"])
        if row["chart_label_en"] == "Final withholding tax":
            adjustment = -capital_adjustment - running
        else:
            adjustment = -round(capital_adjustment * chart_value / capital_total)
            running += adjustment
        new_value = chart_value + adjustment
        row["chart_revenue_eur"] = str(new_value)
        row["chart_revenue_pct_gdp"] = pct(new_value)
        row["netted_adjustments_eur"] = str(adjustment)
        row["netting_notes"] = "BMF cash receipts identify this amount as the solidarity surcharge on capital-income taxes; it is shown separately."


def sort_key(row):
    layer_order = {
        "Employment": 1,
        "Goods/services": 2,
        "Business": 3,
        "Land": 4,
        "Wealth": 5,
        "Environmental": 6,
        "Other": 7,
    }
    return (
        layer_order[row["layer_1"]],
        row["layer_2"],
        row["layer_3"],
        -int(row["chart_revenue_eur"] or 0),
        row["chart_label_en"],
    )


def main():
    euro_rows, euro_by_name = read_eurostat()
    _, bmf_by_name = read_bmf()
    solidarity = solidarity_components()

    rows = []
    rows.extend(make_simple_entries(euro_by_name, solidarity))
    split_rows = make_bmf_split_entries(euro_by_name, bmf_by_name)
    apply_capital_solidarity_adjustment(split_rows, solidarity)
    rows.extend(split_rows)
    rows.append(make_solidarity_entry(euro_by_name, solidarity))
    rows.sort(key=sort_key)

    source_total = sum(int(row["revenue_eur"]) for row in euro_rows)
    chart_total = sum(int(row["chart_revenue_eur"]) for row in rows if row["chart_include"] == "yes")
    if chart_total != source_total:
        raise SystemExit(f"Layered chart total mismatch: {chart_total} vs source {source_total}")

    with OUTPUT.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=FIELDNAMES)
        writer.writeheader()
        writer.writerows(rows)

    print(f"Wrote {OUTPUT}")
    print(f"Rows: {len(rows)}")
    print(f"Chart total: EUR {chart_total:,}")
    print(f"Chart % GDP: {chart_total / GDP_EUR * 100:.6f}%")


if __name__ == "__main__":
    main()
