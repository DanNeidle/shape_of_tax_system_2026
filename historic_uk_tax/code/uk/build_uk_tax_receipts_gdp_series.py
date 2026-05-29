#!/usr/bin/env python3
"""Build a long-run UK tax/public-receipts share-of-GDP series."""

from __future__ import annotations

import csv
import shutil
import tempfile
import urllib.request
from pathlib import Path

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_OUTPUT = PROJECT_ROOT / "code/data/uk_tax/uk_tax_receipts_gdp_share.tsv"
OBR_HISTORICAL_PUBLIC_FINANCES_URL = (
    "https://obr.uk/docs/dlm_uploads/Historical-public-finances-database.xlsx"
)

# The OBR historical workbook stops at 2022-23. This is the same OBR/ONS current
# National Accounts taxes control already used in the current UK sunburst chart.
LATEST_NATIONAL_ACCOUNTS_TAXES_POINT = {
    "year_label": "2024-25",
    "date": "2024-07-01",
    "pct_gdp": "34.536",
    "measure": "National Accounts taxes",
    "source": "OBR March 2026 EFO Table A.5 / nominal GDP Table A.3",
    "confidence": "high",
    "note": "Current chart control total: £1,013,286m National Accounts taxes / £2,934,021m nominal GDP.",
}


def download_workbook(url: str) -> Path:
    tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    tmp.close()
    request = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(request) as response, Path(tmp.name).open("wb") as handle:
        shutil.copyfileobj(response, handle)
    return Path(tmp.name)


def fiscal_year_start(label: str) -> int:
    return int(label.split("-")[0])


def midpoint_date(label: str) -> str:
    return f"{fiscal_year_start(label):04d}-07-01"


def value_is_number(value: object) -> bool:
    return isinstance(value, (int, float))


def build_rows(workbook_path: Path) -> list[dict[str, str]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    sheet = workbook["Receipts (per cent of GDP)"]

    rows: list[dict[str, str]] = []
    for row in sheet.iter_rows(min_row=7, max_col=26, values_only=True):
        year_label = row[0]
        if not isinstance(year_label, str) or "-" not in year_label:
            continue

        start_year = fiscal_year_start(year_label)
        measure = ""
        pct_gdp = None
        confidence = "high"
        note = ""

        national_accounts_taxes = row[25]
        public_sector_current_receipts = row[24]
        central_government_receipts = row[21]

        if start_year >= 1946 and value_is_number(national_accounts_taxes):
            pct_gdp = national_accounts_taxes
            measure = "National Accounts taxes"
            note = "Direct OBR historical public finances National Accounts taxes line."
        elif start_year >= 1900 and value_is_number(public_sector_current_receipts):
            pct_gdp = public_sector_current_receipts
            measure = "Public sector current receipts"
            confidence = "medium"
            note = "Used before the OBR National Accounts taxes line begins; includes some non-tax current receipts."
        elif value_is_number(central_government_receipts):
            pct_gdp = central_government_receipts
            measure = "Central government receipts"
            confidence = "medium"
            note = "Used for 1700-1899 because the OBR workbook has no public-sector or National Accounts taxes total for these years."

        if pct_gdp is None:
            continue

        rows.append(
            {
                "year_label": year_label,
                "date": midpoint_date(year_label),
                "pct_gdp": f"{float(pct_gdp):.6f}",
                "measure": measure,
                "source": "OBR historical public finances database",
                "confidence": confidence,
                "note": note,
            }
        )

    rows.append(LATEST_NATIONAL_ACCOUNTS_TAXES_POINT)
    return sorted(rows, key=lambda row: row["date"])


def write_rows(rows: list[dict[str, str]], output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = ["year_label", "date", "pct_gdp", "measure", "source", "confidence", "note"]
    with output.open("w", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, delimiter="\t", lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    workbook_path = download_workbook(OBR_HISTORICAL_PUBLIC_FINANCES_URL)
    try:
        rows = build_rows(workbook_path)
    finally:
        workbook_path.unlink(missing_ok=True)

    write_rows(rows, DEFAULT_OUTPUT)
    print(f"Wrote {DEFAULT_OUTPUT}")
    print(f"Rows: {len(rows)}")
    print(f"First point: {rows[0]['year_label']} = {rows[0]['pct_gdp']}%")
    print(f"Last point: {rows[-1]['year_label']} = {rows[-1]['pct_gdp']}%")


if __name__ == "__main__":
    main()
